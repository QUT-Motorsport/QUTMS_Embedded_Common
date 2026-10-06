"""Lints every DBC here and writes out/: merged bus DBCs, the MoTeC import, the spreadsheet and the board C.

    python tools/build.py            lint, then write out/
    python tools/build.py --check    lint only, non-zero exit on errors
"""
import argparse
import itertools
import math
import os
import re
import sys
from collections import Counter, defaultdict
from fractions import Fraction

import cantools

sys.path.insert(0, os.path.dirname(__file__))
from dbcwriter import ATTR_DEFS, SEND_TYPES, STATUSES, VARIANTS, Msg, Sig, render  # noqa: E402
from qutms_id import PRIO, SRC, SRC_NODE, TYPE, parse  # noqa: E402

ROOT = os.path.normpath(os.path.join(os.path.dirname(__file__), ".."))
OUT = os.path.join(ROOT, "out")
BUS_FILES = {
    "A": ["qev6_can_a.dbc", "vendor/dti_hv500.dbc", "vendor/sevcon.dbc", "vendor/sbg_ecan.dbc", "gen/bmu_cells.dbc"],
    "B": ["qev6_can_b.dbc"],
}
VENDOR = ("vendor/",)
# A Qev6 BMU or the comp25 one, DTIs or the Sevcon backups: one of each is on the bus at a time.
COMBOS = [("bmu_comp25", "inv_dti"), ("bmu_qev6", "inv_dti"), ("bmu_comp25", "inv_sevcon"), ("bmu_qev6", "inv_sevcon")]
MOTEC_RES = (1, 0.1, 0.01, 0.001)  # channel resolutions a scale may land on
MOTEC_CAN_BUDGET = 200  # of the C125's ~300, leaving room for its own pots, wheel speeds and maths
# both BMU firmwares log to the same channel names so the dash maths doesn't care, which means
# only one fits in the import. Swap after comp
MOTEC_BMU = "bmu_comp25"
MOTEC_PORT = {"A": "CAN1", "B": "CAN2"}


class Rec:
    """One message on one bus, with its attributes resolved."""

    def __init__(self, bus, path, m):
        self.bus, self.path, self.m = bus, path, m
        a = {k: v.value for k, v in (m.dbc.attributes.items() if m.dbc else [])}
        self.cycle = int(a.get("GenMsgCycleTime", 0) or 0)
        self.send_type = a.get("SendType", "Cyclic")
        self.rx_timeout = int(a.get("RxTimeoutMs", 0) or 0)
        self.motec = a.get("MoTeC", "No") == "Yes"
        self.status = a.get("Status", "Current")
        self.variant = a.get("Variant", "")
        self.vendor = path.startswith(VENDOR)
        self.sender = (m.senders or ["?"])[0]

    @property
    def key(self):
        return (self.m.is_extended_frame, self.m.frame_id)

    @property
    def hex_id(self):
        return f"0x{self.m.frame_id:08X}" if self.m.is_extended_frame else f"0x{self.m.frame_id:03X}"

    def logged(self):
        return [s for s in self.m.signals if self.motec and sig_logged(s)]

    def rate(self):
        return 1000.0 / self.cycle if self.cycle and self.send_type in ("Cyclic", "Burst") else 0.0

    def receivers(self):
        rx = []
        for s in self.m.signals:
            rx += [r for r in s.receivers if r not in rx and r != "Vector__XXX"]
        return rx

    def on_bus(self, combo):
        return self.status != "Legacy" and (self.variant == "" or self.variant in combo)


def sig_logged(s):
    a = s.dbc.attributes if s.dbc else {}
    return not ("MoTeCLog" in a and a["MoTeCLog"].value == "No")


def sig_group(s):
    a = s.dbc.attributes if s.dbc else {}
    return a["MoTeCGroup"].value if "MoTeCGroup" in a else ""


class Channel:
    """One C125 channel: a logged signal, or a group of flag bits logged as one word."""

    def __init__(self, s, name=None, start=None, length=None, bits=None):
        self.sig, self.name = s, name or s.name
        self.start = s.start if start is None else start
        self.length = s.length if length is None else length
        self.byte_order, self.is_float = s.byte_order, s.is_float
        self.is_signed = s.is_signed if bits is None else False
        self.scale, self.offset = (s.scale, s.offset) if bits is None else (1, 0)
        self.unit = (s.unit or "") if bits is None else ""
        self.choices = (s.choices or {}) if bits is None else {}
        self.bits = bits  # [(bit in the word, signal name)] for a packed group

    @property
    def minimum(self):
        return self.sig.minimum if self.bits is None else 0

    @property
    def maximum(self):
        return self.sig.maximum if self.bits is None else (1 << self.length) - 1


def motec_channels(r):
    """Flags in the same MoTeCGroup become one word: the C125 tops out around 300 channels all up."""
    out, groups = [], defaultdict(list)
    for s in r.logged():
        (groups[sig_group(s)].append(s) if sig_group(s) else out.append(Channel(s)))
    for g, ss in groups.items():
        if ss[0].byte_order == "little_endian":
            lo, hi = min(s.start for s in ss), max(s.start + s.length - 1 for s in ss)
            bits = [(s.start - lo, s.name) for s in ss]
        else:  # whole bytes, MSB first, so the word reads the same way the DTI sends it
            first, last = min(byte_span(s)[0] for s in ss), max(byte_span(s)[1] for s in ss)
            lo, hi = first * 8 + 7, None
            bits = [((last - s.start // 8) * 8 + s.start % 8, s.name) for s in ss]
        length = (hi - lo + 1) if hi is not None else (last - first + 1) * 8
        out.append(Channel(ss[0], name=g, start=lo, length=length, bits=sorted(bits)))
    return sorted(out, key=lambda c: byte_span(c)[0] * 8 + (c.start % 8))


def byte_span(s):
    """First and last byte a signal touches."""
    if s.byte_order == "little_endian":
        return s.start // 8, (s.start + s.length - 1) // 8
    # Motorola sawtooth: walk from the MSB down.
    first = s.start // 8
    bits_in_first = (s.start % 8) + 1
    rest = max(0, s.length - bits_in_first)
    return first, first + math.ceil(rest / 8)


def motec_scale(s):
    """(resolution, multiplier, divisor, adder) that the C125 can do exactly, or None."""
    for r in MOTEC_RES:
        f = Fraction(s.scale).limit_denominator(10 ** 9) / Fraction(r).limit_denominator(1000)
        o = Fraction(s.offset).limit_denominator(10 ** 9) / Fraction(r).limit_denominator(1000)
        if f.numerator <= 32767 and f.denominator <= 32767 and o.denominator == 1 and abs(o) <= 32767:
            return r, f.numerator, f.denominator, int(o)
    return None


class Lint:
    def __init__(self):
        self.items = []

    def err(self, where, msg):
        self.items.append(("ERROR", where, msg))

    def warn(self, where, msg):
        self.items.append(("WARN", where, msg))

    @property
    def errors(self):
        return [i for i in self.items if i[0] == "ERROR"]


def load():
    recs, lint, nodes = [], Lint(), defaultdict(list)
    want = [l for l in render([], []).splitlines() if l.startswith(("BA_DEF_", "BA_DEF_DEF_"))]
    for bus, files in BUS_FILES.items():
        for f in files:
            path = os.path.join(ROOT, f)
            if not os.path.exists(path):
                lint.err(f, "missing (run the generators?)")
                continue
            text = open(path, encoding="utf-8").read()
            have = [l for l in text.splitlines() if l.startswith(("BA_DEF_", "BA_DEF_DEF_"))]
            if have != want:
                lint.err(f, "attribute definitions differ from tools/dbcwriter.py ATTR_DEFS")
            try:
                db = cantools.database.load_string(text, strict=True)
            except Exception as e:  # cantools raises a few different things
                lint.err(f, f"cantools can't load it: {e}")
                continue
            for n in db.nodes:
                if n.name not in nodes[bus]:
                    nodes[bus].append(n.name)
            recs += [Rec(bus, f, m) for m in db.messages]
    return recs, lint, nodes


def check(recs, lint):
    names = defaultdict(list)
    for r in recs:
        w = f"{r.bus}:{r.m.name}"
        names[r.m.name].append(r)
        if r.status not in STATUSES:
            lint.err(w, f"Status '{r.status}' isn't one of {STATUSES}")
        if r.send_type not in SEND_TYPES:
            lint.err(w, f"SendType '{r.send_type}' isn't one of {SEND_TYPES}")
        if r.variant not in VARIANTS:
            lint.err(w, f"Variant '{r.variant}' isn't one of {VARIANTS}")
        if r.send_type in ("Cyclic", "Burst") and not r.cycle:
            lint.err(w, "cyclic but no GenMsgCycleTime")
        if r.rx_timeout and r.cycle and r.send_type == "Cyclic" and r.rx_timeout < 2 * r.cycle:
            lint.warn(w, f"RxTimeoutMs {r.rx_timeout} is under 2x the {r.cycle} ms period")
        ext, fid = r.key
        if not r.vendor and r.variant != "bmu_comp25" and r.status != "Legacy" and ext and r.sender in set(SRC_NODE.values()) | {"TOOL"}:
            f = parse(fid)
            owner = next((n for s, n in SRC_NODE.items() if s == f["src_name"]), None)
            if f["type_name"] != "RECEIVE" and owner != r.sender:
                lint.err(w, f"ID {r.hex_id} says source {f['src_name']} but the sender is {r.sender}")
            if f["fd"] or f["autonomous"]:
                lint.err(w, f"ID {r.hex_id} sets the FD or autonomous bit")
            if (f["type_name"] == "HEARTBEAT") != (f["prio_name"] == "HEARTBEAT"):
                lint.err(w, f"ID {r.hex_id}: heartbeat type and heartbeat priority should go together")
            if (f["type_name"] == "ERROR") != (f["prio_name"] == "ERROR"):
                lint.err(w, f"ID {r.hex_id}: error type and error priority should go together")
        if not r.vendor and ext and fid <= 0xFFFF:
            lint.err(w, f"ID {r.hex_id} sits in the DTI range (packet << 8 | node)")
        if ext and 0x1CF80000 <= fid <= 0x1CF8FFFF:
            lint.err(w, f"ID {r.hex_id} is in MoTeC's reserved 0x1CF8xxxx device-discovery range")
        if r.bus == "A" and not ext and (fid == 0 or 0x200 <= fid <= 0x20F or 0x280 <= fid <= 0x28F):
            (lint.warn if r.vendor else lint.err)(w, f"ID {r.hex_id} is in the C125's reserved 11-bit range on CAN1")
        if r.motec:
            check_motec(r, lint, w)
    for name, rs in names.items():
        if len(rs) > 1:
            same = all(x.key == rs[0].key and x.m.length == rs[0].m.length for x in rs)
            if not (same and len({x.bus for x in rs}) == len(rs)):
                lint.err(name, "message name used more than once")
    from cantools.database.can.c_source import camel_to_snake_case
    for bus in BUS_FILES:
        snake = {}
        for r in recs:
            if r.bus == bus:
                s = camel_to_snake_case(r.m.name)
                if s in snake and snake[s] != r.m.name:
                    lint.err(f"{bus}:{r.m.name}", f"generated C name '{s}' clashes with {snake[s]}")
                snake[s] = r.m.name
        for combo in COMBOS:
            seen = {}
            for r in recs:
                if r.bus == bus and r.on_bus(combo):
                    if r.key in seen:
                        lint.err(f"{bus}:{r.m.name}", f"ID {r.hex_id} also used by {seen[r.key]} ({'+'.join(combo)})")
                    seen[r.key] = r.m.name
    for combo in COMBOS:
        seen = {}
        for r in recs:
            if r.on_bus(combo) and r.motec:
                for c in motec_channels(r):
                    if c.name in seen and seen[c.name] != r.m.name:
                        lint.err(c.name, f"MoTeC channel name used by {seen[c.name]} and {r.m.name} ({'+'.join(combo)})")
                    seen[c.name] = r.m.name
        if len(seen) > MOTEC_CAN_BUDGET:
            lint.warn("MoTeC", f"{len(seen)} CAN channels for {'+'.join(combo)}: the C125 takes 'over 300' all up, "
                               f"including its own inputs and maths, so keep CAN under {MOTEC_CAN_BUDGET}")
    seen = {}
    for r in motec_import(recs):
        for c in motec_channels(r):
            if c.name in seen and seen[c.name] != r.m.name:
                lint.err(c.name, f"MoTeC channel name used by {seen[c.name]} and {r.m.name} in qev6_motec.dbc")
            seen[c.name] = r.m.name


def motec_import(recs):
    """Everything the C125 could log off either bus, one file. Both inverter options are in, they don't clash."""
    order = lambda r: (r.bus, r.m.is_extended_frame, r.m.frame_id)
    return sorted((r for r in recs if r.motec and r.status != "Legacy" and
                   (r.variant in ("", MOTEC_BMU) or r.variant.startswith("inv_"))), key=order)


def bitset(c):
    """Frame bits a signal or channel covers, in DBC numbering."""
    if c.byte_order == "little_endian":
        return set(range(c.start, c.start + c.length))
    out, b = set(), c.start
    for _ in range(c.length):
        out.add(b)
        b = b - 1 if b % 8 else b + 15  # Motorola sawtooth: bit 0 of a byte runs on to bit 7 of the next
    return out


def check_motec(r, lint, w):
    logged = motec_channels(r)
    for i, c in enumerate(logged):
        for d in logged[i + 1:]:
            if bitset(c) & bitset(d):
                lint.err(f"{w}.{c.name}", f"overlaps {d.name}; pack the group into bytes nothing else uses")
    if len({s.byte_order for s in r.m.signals}) > 1:
        lint.err(w, "mixes byte orders; the C125 takes one per message")
    if r.rate() >= 500:
        lint.err(w, f"{r.rate():.0f} Hz is past the C125's 500 Hz fast receive")
    for s in logged:
        sw = f"{w}.{s.name}"
        a, b = byte_span(s)
        span = b - a + 1
        window = next((x for x in (1, 2, 4) if x >= span), None)
        if s.length > 32 or s.is_float:
            lint.err(sw, "over 32 bits or a float, the C125 can't read it")
        elif window is None or a + window > 8:
            lint.err(sw, f"bytes {a}-{b} don't fit a 1/2/4-byte window inside the frame")
        if motec_scale(s) is None:
            lint.warn(sw, f"scale {s.scale}/offset {s.offset} has no exact multiplier/divisor <= 32767; the import will round")
        if s.is_signed and s.length not in (16, 32):
            lint.warn(sw, f"signed {s.length}-bit: the C125 'Signed' tick is only confirmed for 16/32-bit fields")
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*", s.name) or s.name.upper() in ("VERSION", "BS_", "BU_"):
            lint.err(sw, "not a plain identifier, the DBC import will choke")
        if len(s.name) > 32:
            lint.warn(sw, "long channel name, Dash Manager may truncate it")


def to_msg(r, signals=None, motec_only=False):
    """cantools message back into the writer's objects."""
    sigs = []
    for s in (signals if signals is not None else r.m.signals):
        sigs.append(Sig(s.name, s.start, s.length, le=s.byte_order == "little_endian", signed=s.is_signed,
                        factor=s.scale, offset=s.offset, lo=s.minimum, hi=s.maximum, unit=s.unit or "",
                        values={int(k): str(v) for k, v in (s.choices or {}).items()},
                        comment="" if motec_only else (s.comment or ""), motec=sig_logged(s),
                        receivers=[x for x in s.receivers if x != "C125"], group=sig_group(s)))
    return Msg(r.m.name, r.m.frame_id, r.m.length, r.sender, sigs, ext=r.m.is_extended_frame,
               cycle_ms=r.cycle, send_type=r.send_type, rx_timeout_ms=0 if motec_only else r.rx_timeout,
               motec=r.motec, status="Current" if motec_only else r.status, variant="" if motec_only else r.variant,
               comment="" if motec_only else (r.m.comment or ""))


def motec_msg(r, name=None):
    """What the C125 imports: logged channels only, flag groups as words with the bit map in the comment."""
    sigs = [Sig(c.name, c.start, c.length, le=c.byte_order == "little_endian", signed=c.is_signed,
                factor=c.scale, offset=c.offset, lo=c.minimum, hi=c.maximum, unit=c.unit,
                values={int(k): str(v) for k, v in c.choices.items()},
                comment=", ".join(f"b{b} {n}" for b, n in c.bits) if c.bits else "")
            for c in motec_channels(r)]
    return Msg(name or r.m.name, r.m.frame_id, r.m.length, r.sender, sigs, ext=r.m.is_extended_frame,
               cycle_ms=r.cycle, send_type=r.send_type, motec=True)


def write_outputs(recs, nodes):
    order = lambda r: (r.m.is_extended_frame, r.m.frame_id)
    for bus in BUS_FILES:
        rs = sorted((r for r in recs if r.bus == bus), key=order)
        with open(os.path.join(OUT, f"qev6_can_{bus.lower()}.dbc"), "w", encoding="utf-8", newline="\n") as f:
            f.write(render([to_msg(r) for r in rs], nodes[bus], db_name=f"qev6_can_{bus.lower()}"))
    # the same file goes on both ports, the CAN1_/CAN2_ on the front says which one keeps it
    msgs = [motec_msg(r, f"{MOTEC_PORT[r.bus]}_{r.m.name}") for r in motec_import(recs)]
    used = sorted({m.sender for m in msgs if m.sender})
    with open(os.path.join(OUT, "qev6_motec.dbc"), "w", encoding="utf-8", newline="\n") as f:
        f.write(render(msgs, used, db_name="qev6_motec"))


def write_xlsx(recs):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as col

    fields = [("Priority", range(28, 26, -1), "E06666", "F4CCCC"), ("Source ID", range(26, 21, -1), "6FA8DC", "CFE2F3"),
              ("A", [21], "F6B26B", "FCE5CD"), ("Type", range(20, 17, -1), "8E7CC3", "D9D2E9"),
              ("FD", [17], "76A5AF", "D0E0E3"), ("Extra", range(16, 3, -1), "D9D9D9", "F3F3F3"),
              ("Board No.", range(3, -1, -1), "93C47D", "D9EAD3")]
    bit_field = {b: f for f in fields for b in f[1]}
    groups = [("SHUTDOWN", "Shutdown"), ("BMU", "BMU"), ("ECU", "ECU"), ("CHARGE_CONTROL", "Charge Control"),
              ("ACM", "ACM"), ("SW_MOTEC", "SW (C125)"), ("ROS", "ROS")]
    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    mid = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")
    fill = lambda c: PatternFill("solid", fgColor=c)
    font = Font(name="Arial", size=9)
    bold = Font(name="Arial", size=9, bold=True)

    def put(ws, r, c, v=None, f=None, fnt=font, al=mid):
        cell = ws.cell(row=r, column=c, value=v)
        cell.border, cell.font, cell.alignment = box, fnt, al
        if f:
            cell.fill = fill(f)
        return cell

    def span(ws, r1, c1, r2, c2, v, f=None, fnt=bold, al=mid):
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                put(ws, r, c, None, f, fnt, al)
        ws.cell(row=r1, column=c1, value=v)
        if (r1, c1) != (r2, c2):
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

    def bit_header(ws, row, c0, top=True):
        # row: field names, row + 1: bit numbers. Optional "CAN ID BITS" band above.
        if top:
            span(ws, row - 1, c0, row - 1, c0 + 28, "CAN ID BITS")
        c = c0
        for name, bits, hard, _ in fields:
            n = len(list(bits))
            span(ws, row, c, row, c + n - 1, name, hard)
            for b in bits:
                put(ws, row + 1, c, b, hard, bold)
                c += 1

    def bits_row(ws, r, c0, frame_id, tint):
        for i, b in enumerate(range(28, -1, -1)):
            put(ws, r, c0 + i, (frame_id >> b) & 1, bit_field[b][3] if tint else None)

    seen, scheme, other = set(), [], []
    for r in sorted(recs, key=lambda r: (r.bus, r.m.is_extended_frame, r.m.frame_id)):
        key = (r.m.name, r.m.frame_id)
        if key in seen:
            continue
        seen.add(key)
        f = parse(r.m.frame_id) if r.m.is_extended_frame else None
        if (r.vendor or not r.m.is_extended_frame or r.status == "Legacy" or r.variant == "bmu_comp25"
                or f["src_name"] is None):
            other.append(r)
        else:
            scheme.append((f["src_name"], r))

    wb = Workbook()

    # overview: GROUP, MESSAGE, the bits. Cell families show their first CMU only
    ws = wb.active
    ws.title = "CAN IDs"
    span(ws, 1, 1, 3, 1, "GROUP")
    span(ws, 1, 2, 3, 2, "MESSAGE")
    bit_header(ws, 2, 3)
    row = 4
    for src, label in groups:
        rs = sorted((r for s, r in scheme if s == src and not re.search(r"_N[1-9]\d*(_|$)", r.m.name)),
                    key=lambda r: r.m.frame_id)
        if not rs:
            continue
        span(ws, row, 1, row + len(rs) - 1, 1, label, fnt=font)
        for r in rs:
            put(ws, row, 2, r.m.name, al=left)
            bits_row(ws, row, 3, r.m.frame_id, False)
            row += 1
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 16, 34
    for c in range(3, 32):
        ws.column_dimensions[col(c)].width = 3.2
    ws.freeze_panes = "C4"

    # what each bit pattern means, and everything that isn't on the scheme
    ws = wb.create_sheet("ID Aspects")
    span(ws, 1, 1, 2, 1, "ID/ASPECT", "000000", Font(name="Arial", size=11, bold=True, color="FFFFFF"))
    bit_header(ws, 1, 2, top=False)
    aspects = [("Priority", n.title(), v) for n, v in sorted(PRIO.items(), key=lambda kv: kv[1])]
    aspects += [("Source ID", dict(groups).get(n, n), v) for n, v in sorted(SRC.items(), key=lambda kv: kv[1])]
    aspects += [("A", "EV", 0), ("A", "AV", 1)]
    aspects += [("Type", n if n == "OBJ_DICT" else n.title(), v) for n, v in sorted(TYPE.items(), key=lambda kv: kv[1])]
    aspects += [("FD", "Non FD", 0), ("FD", "FD", 1), ("Extra", "Miscellaneous", 0), ("Board No.", "Board Dependant", 0)]
    row = 3
    for fname, label, val in aspects:
        name, bits, hard, tint = next(f for f in fields if f[0] == fname)
        bits = list(bits)
        put(ws, row, 1, label, None if fname == "Extra" else hard, al=left)
        for i, b in enumerate(range(28, -1, -1)):
            if b in bits:
                put(ws, row, 2 + i, (val >> (b - bits[-1])) & 1, tint)
            else:
                put(ws, row, 2 + i)
        row += 1
    row += 1
    span(ws, row, 1, row, 30, "Other Messages", "000000", Font(name="Arial", size=11, bold=True, color="FFFFFF"))
    row += 1
    put(ws, row, 1, "Message", "999999")
    span(ws, row, 2, row, 30, "ID", "999999", fnt=font)
    row += 1
    fams = {}
    for r in other:
        k = re.sub(r"_N_?\d+", "_N*", re.sub(r"_P_?\d+$", "_P*", r.m.name)) if re.search(r"_N_?\d+", r.m.name) else r.m.name
        fams.setdefault(k, []).append(r)
    for k, rs in fams.items():
        ids = sorted(r.m.frame_id for r in rs)
        hexid = lambda i, ext: f"0x{i:08X}" if ext else f"0x{i:X}"
        ext = rs[0].m.is_extended_frame
        txt = hexid(ids[0], ext) if len(ids) == 1 else f"{hexid(ids[0], ext)} to {hexid(ids[-1], ext)} ({len(ids)} frames)"
        put(ws, row, 1, k, "CCCCCC", al=left)
        span(ws, row, 2, row, 30, txt, "CCCCCC", fnt=font)
        row += 1
    ws.column_dimensions["A"].width = 30
    for c in range(2, 31):
        ws.column_dimensions[col(c)].width = 3.2

    # one sheet per board: MESSAGE, the bits, then the ID three ways
    for src, label in groups:
        rs = sorted((r for s, r in scheme if s == src), key=lambda r: r.m.frame_id)
        if not rs:
            continue
        ws = wb.create_sheet(label)
        span(ws, 1, 1, 3, 1, "MESSAGE")
        bit_header(ws, 2, 2)
        span(ws, 1, 31, 1, 33, "CAN ID")
        for c, h in ((31, "Binary"), (32, "Decimal"), (33, "Hex")):
            span(ws, 2, c, 3, c, h)
        for i, r in enumerate(rs, 4):
            put(ws, i, 1, r.m.name, al=left)
            bits_row(ws, i, 2, r.m.frame_id, True)
            put(ws, i, 31, format(r.m.frame_id, "029b"))
            put(ws, i, 32, r.m.frame_id)
            put(ws, i, 33, f"0x{r.m.frame_id:08X}")
        ws.column_dimensions["A"].width = 34
        for c in range(2, 31):
            ws.column_dimensions[col(c)].width = 3.2
        for c, w in ((31, 32), (32, 12), (33, 13)):
            ws.column_dimensions[col(c)].width = w
        ws.freeze_panes = "B4"

    # the dropdown lists
    ws = wb.create_sheet("Lists")
    lists = [(2, [n.title() for n, _ in sorted(PRIO.items(), key=lambda kv: kv[1])], "E06666"),
             (4, [dict(groups).get(n, n) for n, _ in sorted(SRC.items(), key=lambda kv: kv[1])], "6FA8DC"),
             (6, ["EV", "AV"], "F6B26B"),
             (8, [n if n == "OBJ_DICT" else n.title() for n, _ in sorted(TYPE.items(), key=lambda kv: kv[1])], "8E7CC3"),
             (10, ["Non FD", "FD"], "76A5AF")]
    for c, items, hard in lists:
        for i, v in enumerate(items, 2):
            put(ws, i, c, v, hard, al=left)
        ws.column_dimensions[col(c)].width = 16
    put(ws, 2, 12, "Miscellaneous", al=left)
    put(ws, 5, 12, "Board Dependant", "93C47D", al=left)
    ws.column_dimensions[col(12)].width = 16

    wb.save(os.path.join(OUT, "qev6_can.xlsx"))


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--check", action="store_true", help="lint only")
    a = ap.parse_args()
    recs, lint, nodes = load()
    check(recs, lint)
    for lvl, where, msg in sorted(lint.items):
        print(f"{lvl:5} {where}: {msg}")
    print(f"{len(recs)} messages, {len(lint.errors)} errors, {len(lint.items) - len(lint.errors)} warnings")
    if a.check:
        sys.exit(1 if lint.errors else 0)
    write_outputs(recs, nodes)
    write_xlsx(recs)
    if not lint.errors:
        import gen_c
        for node in gen_c.BOARDS:
            print(node, ", ".join(gen_c.gen(node)))
    print(f"wrote {os.path.relpath(OUT)}/")
    sys.exit(1 if lint.errors else 0)


if __name__ == "__main__":
    main()
